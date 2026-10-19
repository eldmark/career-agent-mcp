"""Application tracker. The CSV is the single source of truth; the xlsx is regenerated from it."""
import csv
from datetime import date, timedelta
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
CSV_PATH = BASE_DIR / "job_applications_tracker.csv"
XLSX_PATH = BASE_DIR / "job_applications_tracker.xlsx"

HEADERS = [
    "ID", "Fecha", "Empresa", "Puesto / Título", "Salario Esperado",
    "Modalidad", "URL de Vacante", "Track CV Usado", "ATS Match",
    "Estado", "Descripción Resumida", "Notas / Siguiente Paso", "Seguimiento",
]
FOLLOWUP_DAYS = 7
CLOSED_STATES = ("rechazado", "oferta", "descartado", "retirado")


def read_rows(path: Path = CSV_PATH) -> list:
    if not path.exists():
        return []
    with open(path, "r", newline="", encoding="utf-8") as f:
        return [{h: row.get(h) or "" for h in HEADERS} for row in csv.DictReader(f)]


def write_rows(rows: list, path: Path = CSV_PATH) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def _next_id(rows: list) -> int:
    ids = [int(r["ID"]) for r in rows if str(r["ID"]).isdigit()]
    return max(ids, default=0) + 1


def _default_followup(status: str, today: date) -> str:
    if status.lower().startswith("aplicado"):
        return (today + timedelta(days=FOLLOWUP_DAYS)).isoformat()
    return ""


def upsert(entry: dict, path: Path = CSV_PATH) -> tuple:
    """Insert a row, or update the existing one with the same URL. Returns (row, created)."""
    rows = read_rows(path)
    today = date.today()
    url = entry.get("URL de Vacante", "")
    existing = next((r for r in rows if url and r["URL de Vacante"] == url), None)
    if existing:
        for key, val in entry.items():
            if val not in ("", None, "N/A"):
                existing[key] = val
        if not entry.get("Seguimiento"):
            existing["Seguimiento"] = _default_followup(existing["Estado"], today) or existing["Seguimiento"]
        row, created = existing, False
    else:
        row = {h: "" for h in HEADERS}
        row.update(entry)
        row["ID"] = _next_id(rows)
        row["Fecha"] = today.isoformat()
        row["Seguimiento"] = entry.get("Seguimiento") or _default_followup(row["Estado"], today)
        rows.append(row)
        created = True
    write_rows(rows, path)
    return row, created


def update_status(app_id: int, status: str, notes: str = "", followup: str = "",
                  path: Path = CSV_PATH) -> dict:
    rows = read_rows(path)
    row = next((r for r in rows if str(r["ID"]) == str(app_id)), None)
    if not row:
        return {"status": "error", "message": f"No application with ID {app_id}"}
    row["Estado"] = status
    if notes:
        row["Notas / Siguiente Paso"] = notes
    if followup:
        row["Seguimiento"] = followup
    elif status.lower().startswith(CLOSED_STATES):
        row["Seguimiento"] = ""
    else:
        row["Seguimiento"] = _default_followup(status, date.today()) or row["Seguimiento"]
    write_rows(rows, path)
    return {"status": "success", "updated": row}


def due_followups(today: date = None, path: Path = CSV_PATH) -> list:
    today = today or date.today()
    due = []
    for r in read_rows(path):
        if r["Estado"].lower().startswith(CLOSED_STATES):
            continue
        try:
            when = date.fromisoformat(r["Seguimiento"])
        except ValueError:
            continue
        if when <= today:
            due.append(r)
    return due


def export_xlsx(rows: list, pdf_resolver=None, path: Path = XLSX_PATH) -> str:
    """Regenerate the xlsx from rows. Returns "" on success or a warning message."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        return "openpyxl not installed: xlsx was not regenerated (CSV is up to date)"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Postulaciones"
    ws.append(HEADERS)
    ws.row_dimensions[1].height = 28
    header_fill = PatternFill(start_color="122644", end_color="122644", fill_type="solid")
    for c in ws[1]:
        c.fill = header_fill
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    side = Side(style="thin", color="D0D5DD")
    border = Border(left=side, right=side, top=side, bottom=side)
    link_font = Font(name="Calibri", size=10, color="004B91", underline="single")
    for r in rows:
        ws.append([int(r["ID"]) if str(r["ID"]).isdigit() else r["ID"]] + [r[h] for h in HEADERS[1:]])
        row_idx = ws.max_row
        ws.row_dimensions[row_idx].height = 24
        for col, h in enumerate(HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border
            wrap = h in ("Descripción Resumida", "Notas / Siguiente Paso")
            center = h in ("ID", "Fecha", "Salario Esperado", "Modalidad", "ATS Match", "Estado", "Seguimiento")
            cell.alignment = Alignment(horizontal="center" if center else None, vertical="center", wrap_text=wrap)
            if h == "URL de Vacante" and r[h].startswith("http"):
                cell.hyperlink, cell.font = r[h], link_font
            if h == "Track CV Usado" and pdf_resolver:
                lang = "es" if "(es)" in r[h].lower() else "en"
                pdf = pdf_resolver(r[h], lang) if r[h] else None
                if pdf:
                    cell.hyperlink, cell.font = f"file://{Path(pdf).resolve()}", link_font
            if h == "Estado":
                cell.font = Font(name="Calibri", size=10, bold=True)
                low = r[h].lower()
                if low.startswith("aplicado"):
                    cell.fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="137333")
                elif low.startswith("entrevista"):
                    cell.fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="1A73E8")
    ws.freeze_panes = "A2"
    try:
        wb.save(path)
    except OSError as e:
        return f"could not write xlsx: {e}"
    return ""
